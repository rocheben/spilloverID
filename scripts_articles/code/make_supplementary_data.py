"""Build the Supplementary Data Excel file (Roche_Supplementary_Data.xlsx).

Requires:
  - case_results.json (produced by run_cases.R)
  - benchmark_results.json (produced by run_benchmark.R)
  - suppl_sensitivity.json (produced by run_suppl.R)
  - suppl_loco.json (produced by run_suppl.R)
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

with open('../data/case_results.json') as f:
    cases = json.load(f)
with open('../data/benchmark_results.json') as f:
    bench = json.load(f)
with open('../data/suppl_sensitivity.json') as f:
    sens = json.load(f)
with open('../data/suppl_loco.json') as f:
    loco = json.load(f)

wb = openpyxl.Workbook()

# Common styles
HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
HEADER_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
NORMAL_FONT = Font(name='Arial', size=10)
ITALIC_FONT = Font(name='Arial', size=10, italic=True)
TITLE_FONT = Font(name='Arial', size=12, bold=True)
BORDER = Border(left=Side(style='thin', color='CCCCCC'),
                right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'),
                bottom=Side(style='thin', color='CCCCCC'))

def style_header_row(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER

def autosize(ws):
    for col in ws.columns:
        try:
            col_letter = col[0].column_letter
        except Exception:
            continue
        max_len = 0
        for cell in col:
            v = cell.value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)


# ======================================================================
# SHEET 1: README
# ======================================================================
ws = wb.active
ws.title = "README"
ws['A1'] = "Roche et al. — Supplementary Data"
ws['A1'].font = Font(name='Arial', size=14, bold=True)
ws['A2'] = ("Supplementary data tables for the manuscript "
            "'A composite plausibility framework for identifying wildlife reservoirs "
            "of zoonotic pathogens from observational data', submitted to Nature Microbiology.")
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[2].height = 50
ws.column_dimensions['A'].width = 100

ws['A4'] = "Sheet"
ws['B4'] = "Content"
style_header_row(ws, 4, 2)
sheet_descriptions = [
    ("README", "This sheet."),
    ("default_weights",
        "Default weights w_c for the 14 criteria. These sum to 1.0 by construction; "
        "criteria returning NA have their weight redistributed proportionally."),
    ("S_Table1_real_cases",
        "Per-host Π_k point estimates and 95% bootstrap percentile intervals for the three real outbreaks."),
    ("S_Table2_criteria_matrix",
        "Normalized per-criterion scores s_{k,c} (in [0, 1]) for every host × criterion pair in the three real cases."),
    ("S_Table3_benchmark_summary",
        "Per-scenario summary statistics for the 11-scenario mechanistic benchmark "
        "(15 replicates per scenario, 165 simulations total)."),
    ("S_Table4_benchmark_perrep",
        "Per-replicate raw outputs for all 165 simulations: scenario, replicate index, "
        "true reservoir, Π scores, rank, margin, max Π."),
    ("S_Table5_sensitivity",
        "Weight perturbation sensitivity analysis (Dirichlet draws, n=200 per case): "
        "rank-1 stability percentage and Π distribution of the true host."),
    ("S_Table6_loco",
        "Leave-one-criterion-out analysis for the three real cases: change in Π of the true host and "
        "resulting rank when each criterion is removed in turn."),
    ("S_Table7_scenario_params",
        "Parameter modifications defining each of the 11 scenarios relative to the canonical S1.")
]
for i, (name, desc) in enumerate(sheet_descriptions):
    r = 5 + i
    ws.cell(row=r, column=1, value=name).font = Font(name='Arial', size=10, bold=True)
    ws.cell(row=r, column=2, value=desc).font = NORMAL_FONT
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[r].height = 35
ws.column_dimensions['A'].width = 28
ws.column_dimensions['B'].width = 90


# ======================================================================
# SHEET 2: default_weights
# ======================================================================
ws = wb.create_sheet("default_weights")
default_weights = [
    ("C1",  "Spatial co-distribution",                            0.04),
    ("C2",  "Reservoir-vs-accidental host shape",                 0.04),
    ("C3",  "Individual exposure association",                    0.06),
    ("C4",  "Dose-response (residence time)",                     0.04),
    ("C5",  "Low cryptic fraction",                               0.03),
    ("C6",  "Exposure × local prevalence interaction",            0.09),
    ("C7",  "Catalytic FOI coherence",                            0.04),
    ("C8",  "Temporal precedence",                                0.03),
    ("C9",  "Mediation",                                          0.06),
    ("C10", "Mechanistic Akaike weights",                         0.07),
    ("C11", "Specificity (negative-control pathogens)",           0.05),
    ("C12", "Inter-epidemic persistence",                         0.16),
    ("C13", "Endemic age-prevalence shape",                       0.14),
    ("C14", "Spatial independence from co-circulating hosts",     0.15),
]
ws['A1'] = "Default criterion weights w_c"
ws['A1'].font = TITLE_FONT
ws['A3'] = "Criterion"; ws['B3'] = "Name"; ws['C3'] = "Default weight"
style_header_row(ws, 3, 3)
for i, (cn, name, w) in enumerate(default_weights):
    r = 4 + i
    ws.cell(row=r, column=1, value=cn)
    ws.cell(row=r, column=2, value=name)
    ws.cell(row=r, column=3, value=w)
    ws.cell(row=r, column=3).number_format = '0.00'
ws.cell(row=18, column=2, value="Sum").font = Font(bold=True)
ws.cell(row=18, column=3, value="=SUM(C4:C17)").number_format = '0.00'
ws.cell(row=18, column=3).font = Font(bold=True)
autosize(ws)


# ======================================================================
# SHEET 3: S_Table1_real_cases (Π estimates)
# ======================================================================
ws = wb.create_sheet("S_Table1_real_cases")
ws['A1'] = "Supplementary Table 1. Composite plausibility scores for the three real outbreaks"
ws['A1'].font = TITLE_FONT
headers = ["Case", "Host", "Π_k", "95% CI low", "95% CI high", "Rank", "True reservoir?"]
for c, h in enumerate(headers, 1):
    ws.cell(row=3, column=c, value=h)
style_header_row(ws, 3, len(headers))
r = 4
for case_key in ['sn', 'nv', 'rv']:
    c = cases[case_key]
    for i, host in enumerate(c['hosts']):
        ws.cell(row=r, column=1, value=c['case'])
        ws.cell(row=r, column=2, value=host)
        ws.cell(row=r, column=3, value=round(c['pi'][i], 3))
        if c.get('ci_low'):
            ws.cell(row=r, column=4, value=round(c['ci_low'][i], 3))
            ws.cell(row=r, column=5, value=round(c['ci_high'][i], 3))
        ws.cell(row=r, column=6, value=c['rank'][i])
        ws.cell(row=r, column=7, value="YES" if host == c['true_host'] else "")
        if host == c['true_host']:
            for col in range(1, 8):
                ws.cell(row=r, column=col).font = Font(bold=True, color='2ca02c'.upper())
        r += 1
autosize(ws)


# ======================================================================
# SHEET 4: S_Table2_criteria_matrix
# ======================================================================
ws = wb.create_sheet("S_Table2_criteria_matrix")
ws["A1"] = "Supplementary Table 2. Per-criterion normalized scores s_{k,c} (in [0, 1])"
ws['A1'].font = TITLE_FONT
crit_names = [f'C{i}' for i in range(1, 15)]
headers = ["Case", "Host", "True reservoir?"] + crit_names
for c, h in enumerate(headers, 1):
    ws.cell(row=3, column=c, value=h)
style_header_row(ws, 3, len(headers))
r = 4
for case_key in ['sn', 'nv', 'rv']:
    c = cases[case_key]
    for i, host in enumerate(c['hosts']):
        ws.cell(row=r, column=1, value=c['case'])
        ws.cell(row=r, column=2, value=host)
        ws.cell(row=r, column=3, value="YES" if host == c['true_host'] else "")
        crit_row = c['criteria'][i]
        for j, cn in enumerate(crit_names):
            v = crit_row.get(cn)
            if v is None:
                ws.cell(row=r, column=4+j, value="NA")
            else:
                ws.cell(row=r, column=4+j, value=round(float(v), 3))
        if host == c['true_host']:
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).font = Font(bold=True, color='2ca02c'.upper())
        r += 1
autosize(ws)


# ======================================================================
# SHEET 5: S_Table3_benchmark_summary
# ======================================================================
ws = wb.create_sheet("S_Table3_benchmark_summary")
ws['A1'] = "Supplementary Table 3. Benchmark summary per scenario"
ws['A1'].font = TITLE_FONT
headers = ["Scenario", "Label", "N replicates", "Top-1 rate (%)", "Top-2 rate (%)",
           "Mean margin", "SD margin", "Margin < 0 (%)", "Median Π true",
           "Median max Π", "Median Π range", "Median observable margin"]
for c, h in enumerate(headers, 1):
    ws.cell(row=3, column=c, value=h)
style_header_row(ws, 3, len(headers))
r = 4
for row in bench['summary']:
    ws.cell(row=r, column=1, value=row['scenario'])
    ws.cell(row=r, column=2, value=row['label'])
    ws.cell(row=r, column=3, value=row['n_rep'])
    ws.cell(row=r, column=4, value=row.get('top1_rate', 'N/A'))
    ws.cell(row=r, column=5, value=row.get('top2_rate', 'N/A'))
    ws.cell(row=r, column=6, value=row.get('mean_margin', 'N/A'))
    ws.cell(row=r, column=7, value=row.get('sd_margin', 'N/A'))
    ws.cell(row=r, column=8, value=row.get('pct_margin_negative', 'N/A'))
    ws.cell(row=r, column=9, value=row.get('median_pi_true', 'N/A'))
    ws.cell(row=r, column=10, value=row['median_max_pi'])
    ws.cell(row=r, column=11, value=row['median_pi_range'])
    ws.cell(row=r, column=12, value=row.get('median_obs_margin', 'N/A'))
    if row['scenario'] == 'S11_no_spillover':
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).font = Font(bold=True, italic=True, color='888888')
    r += 1
autosize(ws)


# ======================================================================
# SHEET 6: S_Table4_benchmark_perrep
# ======================================================================
ws = wb.create_sheet("S_Table4_benchmark_perrep")
ws['A1'] = ("Supplementary Table 4. Per-replicate benchmark outputs "
            f"({len(bench['per_rep'])} simulations)")
ws['A1'].font = TITLE_FONT
headers = ["Scenario", "Replicate", "True reservoir host id", "Π true", "Max Π other",
           "Margin", "Observable margin", "Rank true", "Top-1", "Top-2",
           "Max Π (all hosts)", "Π range"]
for c, h in enumerate(headers, 1):
    ws.cell(row=3, column=c, value=h)
style_header_row(ws, 3, len(headers))
r = 4
for row in bench['per_rep']:
    vals = [row.get('scenario'), row.get('rep'), row.get('k_star'),
            row.get('pi_true'), row.get('pi_max_other'), row.get('margin'),
            row.get('obs_margin'), row.get('rank_true'), row.get('top1'),
            row.get('top2'), row.get('max_pi'), row.get('pi_range')]
    for c, v in enumerate(vals, 1):
        if v is None:
            ws.cell(row=r, column=c, value="")
        elif isinstance(v, float):
            ws.cell(row=r, column=c, value=round(v, 4))
        else:
            ws.cell(row=r, column=c, value=v)
    r += 1
autosize(ws)
ws.freeze_panes = 'A4'


# ======================================================================
# SHEET 7: S_Table5_sensitivity
# ======================================================================
ws = wb.create_sheet("S_Table5_sensitivity")
ws['A1'] = "Supplementary Table 5. Weight perturbation sensitivity"
ws['A1'].font = TITLE_FONT
ws['A2'] = ("For each case, 200 Dirichlet perturbations of the default weight vector "
            "were drawn (α = default_w × 30). The table reports the rank-1 stability and Π distribution.")
ws['A2'].font = ITALIC_FONT
ws.row_dimensions[2].height = 30
ws.column_dimensions['A'].width = 30
ws.merge_cells('A2:G2')
ws['A2'].alignment = Alignment(wrap_text=True)
headers = ["Case", "True host", "Baseline Π", "% rank-1", "% rank in top-2",
           "Π median", "Π 2.5%", "Π 97.5%"]
for c, h in enumerate(headers, 1):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, len(headers))
r = 5
for key in ['sn', 'nv', 'rv']:
    s = sens[key]
    ws.cell(row=r, column=1, value=s['case'])
    ws.cell(row=r, column=2, value=s['true_host'])
    ws.cell(row=r, column=3, value=round(s['baseline_pi'], 3))
    ws.cell(row=r, column=4, value=round(s['pct_rank1'], 1))
    ws.cell(row=r, column=5, value=round(s['pct_rank_top2'], 1))
    ws.cell(row=r, column=6, value=round(s['pi_median'], 3))
    ws.cell(row=r, column=7, value=round(s['pi_2.5'], 3))
    ws.cell(row=r, column=8, value=round(s['pi_97.5'], 3))
    r += 1
autosize(ws)


# ======================================================================
# SHEET 8: S_Table6_loco
# ======================================================================
ws = wb.create_sheet("S_Table6_loco")
ws['A1'] = "Supplementary Table 6. Leave-one-criterion-out (LOCO)"
ws['A1'].font = TITLE_FONT
ws['A2'] = ("Each criterion's weight is set to zero in turn and weights are renormalized over the "
            "remaining 13 criteria. ΔΠ is the change in the true host's Π relative to baseline.")
ws['A2'].font = ITALIC_FONT
ws.row_dimensions[2].height = 30
ws.merge_cells('A2:F2')
ws['A2'].alignment = Alignment(wrap_text=True)
headers = ["Case", "True host", "Criterion removed", "Π true", "Rank true", "ΔΠ from baseline"]
for c, h in enumerate(headers, 1):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, len(headers))
r = 5
for key in ['sn', 'nv', 'rv']:
    l = loco[key]
    for crow in l['rows']:
        ws.cell(row=r, column=1, value=l['case'])
        ws.cell(row=r, column=2, value=l['true_host'])
        ws.cell(row=r, column=3, value=crow['criterion'])
        ws.cell(row=r, column=4, value=round(crow['pi_true'], 4))
        ws.cell(row=r, column=5, value=crow['rank_true'])
        ws.cell(row=r, column=6, value=round(crow['delta_pi'], 4))
        if crow['rank_true'] > 1:
            for col in range(1, 7):
                ws.cell(row=r, column=col).font = Font(bold=True, color='d62728'.upper())
        r += 1
autosize(ws)
ws.freeze_panes = 'A5'


# ======================================================================
# SHEET 9: S_Table7_scenario_params
# ======================================================================
ws = wb.create_sheet("S_Table7_scenario_params")
ws['A1'] = "Supplementary Table 7. Scenario parameter modifications"
ws['A1'].font = TITLE_FONT
ws['A2'] = ("All scenarios start from the canonical S1 parameter set; the third column lists "
            "the parameter modifications applied. K = 4 candidate hosts in all scenarios.")
ws['A2'].font = ITALIC_FONT
ws.row_dimensions[2].height = 30
ws.merge_cells('A2:C2')
ws['A2'].alignment = Alignment(wrap_text=True)
headers = ["Scenario", "Label", "Parameter modifications vs S1 canonical"]
for c, h in enumerate(headers, 1):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, len(headers))
scenarios = [
    ("S1_canonical", "canonical",
     "Baseline. K=4 hosts, 30 sites, 60 humans/site, 40 animals/host/site, 15 years, R0=2.5, α_reservoir=0.05, α_env=0.003, α_crypt=0.0005."),
    ("S2_bridge", "bridge host",
     "Bridge host k=2; α_reservoir=0.035; α_bridge=0.025."),
    ("S3_cryptic", "cryptic",
     "α_reservoir=0.015; α_env=0.025; α_crypt=0.003 (most transmission via environment)."),
    ("S4_confounded", "confounded",
     "Strong anthropization effects: abundance c(1.0,-0.5,0,0.3); exposure 0.8; prevalence 0.7."),
    ("S5_codistributed", "co-distributed",
     "Two hosts share spatial range: abundance c(0.5,0.5,0,-0.2)."),
    ("S6_strong_seasonal", "seasonal forcing",
     "Seasonal amplitude 0.7; env_noise_sd 0.25."),
    ("S7_low_power", "low power",
     "Reduced sampling: 15 sites, 20 humans/site, 15 animals/host/site."),
    ("S8_measurement_error", "measurement error",
     "exposure misclassification 30%; sensitivity 0.80; specificity 0.90."),
    ("S9_weak_signal", "weak signal",
     "α_reservoir=0.010; α_crypt=0.002."),
    ("S10_short_immunity", "short antibody persistence",
     "antibody_persistence=3 years (vs ~999 default)."),
    ("S11_no_spillover", "NEGATIVE CONTROL — no spillover",
     "α_reservoir = α_env = α_bridge = 0; true_reservoir = NA; +2% baseline animal seroprevalence to mimic assay cross-reactivity.")
]
r = 5
for sid, lab, desc in scenarios:
    ws.cell(row=r, column=1, value=sid)
    ws.cell(row=r, column=2, value=lab)
    ws.cell(row=r, column=3, value=desc)
    ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r].height = 35
    if sid == 'S11_no_spillover':
        for col in range(1, 4):
            ws.cell(row=r, column=col).font = Font(bold=True, italic=True, color='888888')
    r += 1
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 80


wb.save('../manuscript/Roche_Supplementary_Data.xlsx')
print("Supplementary data Excel saved.")
